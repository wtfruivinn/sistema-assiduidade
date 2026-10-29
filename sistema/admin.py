from django.contrib import admin
from django.utils import timezone
from django.db import transaction
from django.db.models import Sum
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from .models import (
    CategoriaDocente, Docente, Curso, UC, Turma, Recinto, Sala,
    Estudante, Inscricao, Horario, Aula, ValidacaoAula,
    Presenca, HorasDocente, FolhaPagamento, Auditoria,
    PedidoRevisao, PedidoHorario  # <-- ADICIONADO AQUI
)
from .auditoria import registar_auditoria
 
 
@admin.register(CategoriaDocente)
class CategoriaDocenteAdmin(admin.ModelAdmin):
    list_display = ("designacao", "valor_hora", "data_inicio", "data_fim")
    search_fields = ("designacao",)
 
 
# ============================================================
# GERAR FOLHA DE PAGAMENTO (a partir do Docente)
# ============================================================
 
def _gerar_folha(modeladmin, request, queryset, mes, ano):
    geradas = 0
    ignoradas = 0
    for docente in queryset:
        folha = FolhaPagamento.objects.filter(docente=docente, mes=mes, ano=ano).first()
        if folha and folha.estado != "rascunho":
            ignoradas += 1
            continue
 
        horas = HorasDocente.objects.filter(
            docente=docente, estado="validada", data__month=mes, data__year=ano
        )
        totais = horas.aggregate(h=Sum("horas"), v=Sum("valor_total"))
        folha, _ = FolhaPagamento.objects.update_or_create(
            docente=docente, mes=mes, ano=ano,
            defaults={
                "total_horas": totais["h"] or 0,
                "valor_total": totais["v"] or 0,
                "estado": "rascunho",
            },
        )
        registar_auditoria(
            request, "FolhaPagamento", folha.id, "gerar_folha",
            detalhes=(f"Docente {docente.numero_docente} - {docente.nome}. "
                      f"Período {mes:02d}/{ano}. Horas: {folha.total_horas}. "
                      f"Valor: {folha.valor_total} MT."),
        )
        geradas += 1
    modeladmin.message_user(
        request,
        f"{geradas} folha(s) gerada(s) para {mes:02d}/{ano}. {ignoradas} ignorada(s) (já validadas/pagas).",
    )
 
 
@admin.action(description="📄 Gerar folha de pagamento do MÊS ATUAL")
def gerar_folha_mes_atual(modeladmin, request, queryset):
    hoje = timezone.localdate()
    _gerar_folha(modeladmin, request, queryset, hoje.month, hoje.year)
 
 
@admin.action(description="📄 Gerar folha de pagamento do MÊS ANTERIOR")
def gerar_folha_mes_anterior(modeladmin, request, queryset):
    hoje = timezone.localdate()
    mes = hoje.month - 1 or 12
    ano = hoje.year if hoje.month > 1 else hoje.year - 1
    _gerar_folha(modeladmin, request, queryset, mes, ano)
 
 
@admin.register(Docente)
class DocenteAdmin(admin.ModelAdmin):
    list_display = ("numero_docente", "nome", "categoria", "departamento", "estado")
    list_filter = ("categoria", "departamento", "estado")
    search_fields = ("numero_docente", "nome")
    actions = [gerar_folha_mes_atual, gerar_folha_mes_anterior]
 
 
@admin.register(Curso)
class CursoAdmin(admin.ModelAdmin):
    list_display = ("codigo", "nome", "faculdade", "duracao_anos", "estado")
    list_filter = ("faculdade", "estado")
    search_fields = ("codigo", "nome")
 
 
@admin.register(UC)
class UCAdmin(admin.ModelAdmin):
    list_display = ("codigo", "nome", "curso", "creditos", "carga_horaria")
    list_filter = ("curso",)
    search_fields = ("codigo", "nome")
 
 
@admin.register(Turma)
class TurmaAdmin(admin.ModelAdmin):
    list_display = ("codigo", "uc", "ano_letivo", "semestre", "capacidade", "estado")
    list_filter = ("ano_letivo", "semestre", "estado")
    search_fields = ("codigo",)
 
 
@admin.register(Recinto)
class RecintoAdmin(admin.ModelAdmin):
    list_display = ("nome", "latitude", "longitude", "raio_metros", "estado")
    list_filter = ("estado",)
    search_fields = ("nome",)
 
 
@admin.register(Sala)
class SalaAdmin(admin.ModelAdmin):
    list_display = ("codigo", "edificio", "piso", "recinto", "wifi_ssid", "estado")
    list_filter = ("recinto", "edificio", "estado")
    search_fields = ("codigo", "edificio")
 
 
@admin.register(Estudante)
class EstudanteAdmin(admin.ModelAdmin):
    list_display = ("numero_estudante", "nome", "curso", "ano_ingresso", "estado")
    list_filter = ("curso", "estado")
    search_fields = ("numero_estudante", "nome")
 
 
@admin.register(Inscricao)
class InscricaoAdmin(admin.ModelAdmin):
    list_display = ("estudante", "turma", "data_inscricao", "estado")
    list_filter = ("estado",)
 
 
@admin.register(Horario)
class HorarioAdmin(admin.ModelAdmin):
    list_display = ("turma", "docente", "sala", "dia_semana", "hora_inicio", "hora_fim")
    list_filter = ("dia_semana", "ano_letivo", "semestre")
 
 
@admin.register(Aula)
class AulaAdmin(admin.ModelAdmin):
    list_display = ("turma", "docente", "sala", "data", "hora_inicio", "hora_fim", "estado")
    list_filter = ("estado", "data")
    search_fields = ("turma__codigo", "docente__nome")
 
 
@admin.register(ValidacaoAula)
class ValidacaoAulaAdmin(admin.ModelAdmin):
    list_display = ("aula", "docente", "data_hora", "resultado", "sala_validada", "recinto_validado")
    list_filter = ("resultado", "sala_validada", "recinto_validado")
 
 
@admin.register(Presenca)
class PresencaAdmin(admin.ModelAdmin):
    list_display = ("aula", "estudante", "estado", "data_hora")
    list_filter = ("estado",)
 
 
# ============================================================
# HORAS DOS DOCENTES
# ============================================================
 
@admin.action(description="✅ Aprovar horas pendentes selecionadas")
def aprovar_horas(modeladmin, request, queryset):
    total = 0
    with transaction.atomic():
        for horas in queryset.filter(estado="pendente"):
            horas.estado = "validada"
            horas.save()
            registar_auditoria(
                request, "HorasDocente", horas.id, "aprovar_horas",
                detalhes=(f"Docente {horas.docente.numero_docente}. Aula {horas.aula_id}. "
                          f"{horas.horas}h = {horas.valor_total} MT. pendente -> validada."),
            )
            total += 1
    modeladmin.message_user(request, f"{total} registo(s) de horas aprovado(s).")
 
 
@admin.register(HorasDocente)
class HorasDocenteAdmin(admin.ModelAdmin):
    list_display = ("docente", "aula", "data", "horas", "valor_hora", "valor_total", "estado")
    list_filter = ("estado", "data")
    actions = [aprovar_horas]
 
 
# ============================================================
# FOLHA DE PAGAMENTO
# ============================================================
 
@admin.action(description="✅ Validar folhas selecionadas")
def validar_folhas(modeladmin, request, queryset):
    total = 0
    with transaction.atomic():
        for folha in queryset.filter(estado="rascunho"):
            folha.estado = "validada"
            folha.data_validacao = timezone.now()
            folha.save()
            registar_auditoria(
                request, "FolhaPagamento", folha.id, "validar_folha",
                detalhes=(f"Docente {folha.docente.numero_docente}. Período {folha.mes:02d}/{folha.ano}. "
                          f"{folha.total_horas}h = {folha.valor_total} MT. rascunho -> validada."),
            )
            total += 1
    modeladmin.message_user(request, f"{total} folha(s) validada(s).")
 
 
@admin.action(description="💰 Marcar folhas como PAGAS")
def marcar_folhas_pagas(modeladmin, request, queryset):
    total = 0
    with transaction.atomic():
        for folha in queryset.filter(estado="validada"):
            folha.estado = "paga"
            folha.save()
            horas_pagas = HorasDocente.objects.filter(
                docente=folha.docente, estado="validada",
                data__month=folha.mes, data__year=folha.ano,
            ).update(estado="paga")
            registar_auditoria(
                request, "FolhaPagamento", folha.id, "pagar_folha",
                detalhes=(f"Docente {folha.docente.numero_docente}. Período {folha.mes:02d}/{folha.ano}. "
                          f"{folha.valor_total} MT. validada -> paga. "
                          f"{horas_pagas} registo(s) de horas passaram a pagas."),
            )
            total += 1
    modeladmin.message_user(request, f"{total} folha(s) marcada(s) como paga(s).")
 
 
@admin.action(description="📊 Exportar folhas selecionadas para Excel")
def exportar_folhas_excel(modeladmin, request, queryset):
    wb = Workbook()
    ws = wb.active
    ws.title = "Folha de Pagamento"
 
    ws.append(["Nº Docente", "Docente", "Categoria", "Mês/Ano",
               "Total de horas", "Valor total (MT)", "Estado"])
    for celula in ws[1]:
        celula.font = Font(bold=True)
 
    soma_horas = 0.0
    soma_valor = 0.0
    folhas = queryset.select_related("docente", "docente__categoria").order_by(
        "ano", "mes", "docente__nome"
    )
    ids = []
    for f in folhas:
        ids.append(str(f.id))
        horas = float(f.total_horas)
        valor = float(f.valor_total)
        soma_horas += horas
        soma_valor += valor
        ws.append([
            f.docente.numero_docente,
            f.docente.nome,
            f.docente.categoria.designacao,
            f"{f.mes:02d}/{f.ano}",
            horas,
            valor,
            f.get_estado_display(),
        ])
 
    ws.append([])
    ws.append(["", "", "", "TOTAL", soma_horas, soma_valor, ""])
    for celula in ws[ws.max_row]:
        celula.font = Font(bold=True)
 
    for coluna, largura in zip("ABCDEFG", (14, 30, 26, 12, 16, 18, 12)):
        ws.column_dimensions[coluna].width = largura
 
    registar_auditoria(
        request, "FolhaPagamento", 0, "exportar_folhas_excel",
        detalhes=f"{len(ids)} folha(s) exportada(s). Ids: {', '.join(ids)}. "
                 f"Total: {soma_horas}h = {soma_valor} MT.",
    )
 
    resposta = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    resposta["Content-Disposition"] = 'attachment; filename="folha_pagamento.xlsx"'
    wb.save(resposta)
    return resposta
 
 
@admin.register(FolhaPagamento)
class FolhaPagamentoAdmin(admin.ModelAdmin):
    list_display = ("docente", "mes", "ano", "total_horas", "valor_total", "estado")
    list_filter = ("estado", "ano", "mes")
    actions = [validar_folhas, marcar_folhas_pagas, exportar_folhas_excel]
 
 
# ============================================================
# AUDITORIA — SÓ DE LEITURA
# ============================================================
 
@admin.register(Auditoria)
class AuditoriaAdmin(admin.ModelAdmin):
    list_display = ("data_hora", "utilizador", "entidade", "entidade_id", "operacao", "ip", "localizacao")
    list_filter = ("entidade", "operacao")
    search_fields = ("utilizador__username", "operacao", "detalhes")
    date_hierarchy = "data_hora"
    ordering = ("-data_hora",)
 
    def has_add_permission(self, request):
        return False
 
    def has_change_permission(self, request, obj=None):
        return False
 
    def has_delete_permission(self, request, obj=None):
        return False
 
 
# ============================================================
# AÇÕES DA COORDENAÇÃO — PEDIDOS DE REVISÃO
# ============================================================
 
@admin.action(description="✅ Aprovar pedidos selecionados")
def aprovar_pedidos(modeladmin, request, queryset):
    total = 0
    with transaction.atomic():
        for pedido in queryset.filter(estado="pendente"):
            aula = pedido.aula
            pedido.estado = "aprovado"
            pedido.decidido_por = request.user
            pedido.data_decisao = timezone.now()
            pedido.save()
 
            ValidacaoAula.objects.create(
                aula=aula,
                docente=aula.docente,
                sala_validada=True,
                recinto_validado=True,
                horario_validado=True,
                qr_validado=False,
                resultado="autorizado",
            )
            estado_anterior = aula.estado
            if aula.estado == "bloqueada":
                aula.estado = "em_curso"
                aula.save()
 
            registar_auditoria(
                request, "PedidoRevisao", pedido.id, "aprovar_revisao",
                detalhes=(f"Aula {aula.id}. Docente {aula.docente.numero_docente}. "
                          f"Aula: {estado_anterior} -> {aula.estado}. "
                          f"Validação manual autorizada pela coordenação."),
            )
            total += 1
    modeladmin.message_user(request, f"{total} pedido(s) aprovado(s).")
 
 
@admin.action(description="❌ Rejeitar pedidos selecionados")
def rejeitar_pedidos(modeladmin, request, queryset):
    total = 0
    with transaction.atomic():
        for pedido in queryset.filter(estado="pendente"):
            pedido.estado = "rejeitado"
            pedido.decidido_por = request.user
            pedido.data_decisao = timezone.now()
            pedido.save()
            registar_auditoria(
                request, "PedidoRevisao", pedido.id, "rejeitar_revisao",
                detalhes=f"Aula {pedido.aula_id}. Docente {pedido.docente.numero_docente}.",
            )
            total += 1
    modeladmin.message_user(request, f"{total} pedido(s) rejeitado(s).")
 
 
@admin.register(PedidoRevisao)
class PedidoRevisaoAdmin(admin.ModelAdmin):
    list_display = ("aula", "docente", "estado", "data_pedido", "decidido_por")
    list_filter = ("estado",)
    search_fields = ("docente__nome", "motivo")
    readonly_fields = ("data_pedido",)
    actions = [aprovar_pedidos, rejeitar_pedidos]


# ============================================================
# ADMIN — PEDIDO DE HORÁRIO (NOVO)
# ============================================================
@admin.register(PedidoHorario)
class PedidoHorarioAdmin(admin.ModelAdmin):
    list_display = ("docente", "turma", "sala", "dia_semana", "hora_inicio", "hora_fim", "estado")
    list_filter = ("estado", "dia_semana")
    search_fields = ("docente__nome", "turma__codigo")
    readonly_fields = ("data_pedido", "decidido_por", "data_decisao")